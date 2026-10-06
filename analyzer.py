"""Import, calculation and Excel export for the Allegro analyzer."""
from io import BytesIO, StringIO
import csv
import math
import pandas as pd
from openpyxl.styles import PatternFill, Font

INPUTS = ["Товар", "SKU", "Закупка PLN", "Доставка PLN", "Пошлина %", "Оформление PLN", "Доп. расходы PLN", "Невозмещаемый НДС PLN", "Цена Allegro PLN", "Комиссия %", "Комиссия PLN"]
NUMERIC = INPUTS[2:]


def number(value):
    if value is None or pd.isna(value) or str(value).strip() == "":
        return 0.0
    s = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    result = float(s)
    if not math.isfinite(result) or result < 0:
        raise ValueError("Числа должны быть конечными и неотрицательными")
    return result


def read_upload(data, filename, sheet=0):
    if filename.lower().endswith(".xlsx"):
        return pd.read_excel(BytesIO(data), sheet_name=sheet, dtype=str).fillna("")
    for encoding in ("utf-8-sig", "cp1250", "cp1251"):
        try:
            content = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Не удалось прочитать кодировку CSV")
    try:
        delimiter = csv.Sniffer().sniff(content[:8192], delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";" if ";" in content else ","
    return pd.read_csv(StringIO(content), sep=delimiter, dtype=str, keep_default_na=False)


def calculate(frame, weak_margin=15):
    result = frame.copy()
    for col in NUMERIC:
        result[col] = result[col].map(number)
    if (result["Комиссия %"] > 100).any():
        raise ValueError("Комиссия Allegro не может превышать 100%")
    result["Пошлина PLN"] = (result["Закупка PLN"] + result["Доставка PLN"]) * result["Пошлина %"] / 100
    result["Себестоимость PLN"] = result[["Закупка PLN", "Доставка PLN", "Пошлина PLN", "Оформление PLN", "Доп. расходы PLN", "Невозмещаемый НДС PLN"]].sum(axis=1)
    result["Комиссия всего PLN"] = result["Цена Allegro PLN"] * result["Комиссия %"] / 100 + result["Комиссия PLN"]
    result["Прибыль PLN"] = result["Цена Allegro PLN"] - result["Себестоимость PLN"] - result["Комиссия всего PLN"]
    result["Маржа %"] = result["Прибыль PLN"].div(result["Цена Allegro PLN"].replace(0, float("nan"))) * 100
    result["ROI %"] = result["Прибыль PLN"].div(result["Себестоимость PLN"].replace(0, float("nan"))) * 100
    def status(row):
        if row["Цена Allegro PLN"] == 0:
            return "Без цены"
        if row["Прибыль PLN"] < 0:
            return "Убыточный"
        if row["Прибыль PLN"] == 0 or row["Маржа %"] < weak_margin:
            return "Слабая прибыль"
        return "Выгодный"
    result["Статус"] = result.apply(status, axis=1)
    return result.sort_values("Прибыль PLN", ascending=False, kind="stable").reset_index(drop=True)


COLORS = {"Выгодный": "DCFCE7", "Слабая прибыль": "FEF3C7", "Убыточный": "FEE2E2", "Без цены": "E2E8F0"}


def export_excel(result, weak_margin=15):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        result.to_excel(writer, sheet_name="Анализ", index=False)
        ws = writer.sheets["Анализ"]
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
        for row, (_, record) in zip(ws.iter_rows(min_row=2), result.iterrows()):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=COLORS[record["Статус"]])
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
                # Prevent spreadsheet formulas from imported product text.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(44, max(18, len(str(column[0].value)) + 3))
        notes = pd.DataFrame({"Правила": [
            "Все суммы и результаты относятся к одной единице товара, в PLN.",
            "Пошлина = (закупка + доставка) × ставка пошлины / 100.",
            "Себестоимость = закупка + доставка + пошлина + оформление + дополнительные расходы + невозмещаемый НДС.",
            "Комиссия = цена продажи × комиссия % / 100 + фиксированная комиссия PLN.",
            "Прибыль = цена продажи − себестоимость − комиссия; маржа = прибыль / цена; ROI = прибыль / себестоимость.",
            f"Слабая прибыль: маржа ниже {weak_margin}% или нулевая прибыль. Без цены: цена не задана.",
            "Нулевой знаменатель: показатель не определён. Возмещаемый НДС не включайте в расходы.",
            "Сравнивайте закупку и продажу на одинаковой базе НДС. Это модель прибыльности, не налоговая декларация."
        ]})
        notes.to_excel(writer, sheet_name="Методика", index=False)
        writer.sheets["Методика"].column_dimensions["A"].width = 120
    return output.getvalue()
